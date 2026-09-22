from __future__ import annotations

import argparse
import json
import re
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "openpyxl is required. Install it with: py -m pip install openpyxl"
    ) from exc


VOCAB_SHEET = "词汇"
CHAR_SHEET = "喃字"


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_latin(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().lower())


def split_latin_variants(value: object) -> list[str]:
    # Handles cells such as "mười, mươi" and "thật; thực".
    text = clean_text(value)
    if not text:
        return []
    return [
        part.strip()
        for part in re.split(r"[,，;；]+", text)
        if part.strip()
    ]


def split_nom_variants(value: object) -> list[str]:
    # Handles the alternative styles already present in the workbook:
    # A/B and A（B）.
    text = clean_text(value)
    if not text:
        return []

    if "/" in text:
        return [part.strip() for part in text.split("/") if part.strip()]

    match = re.fullmatch(r"\s*([^（）]+?)\s*（([^（）]+)）\s*", text)
    if match:
        return [
            match.group(1).strip(),
            match.group(2).strip(),
        ]

    return [text]


def header_map(worksheet) -> dict[str, int]:
    headers = [
        clean_text(cell.value)
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
    ]
    return {header: index for index, header in enumerate(headers) if header}


def require_column(columns: dict[str, int], name: str, sheet_name: str) -> int:
    if name not in columns:
        raise SystemExit(
            f'Sheet "{sheet_name}" is missing required column "{name}". '
            f"Available: {list(columns)}"
        )
    return columns[name]


def optional_value(row: tuple[object, ...], index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return clean_text(row[index])


def append_source(metadata: dict[str, object], source: dict[str, object]) -> None:
    sources = metadata.setdefault("sources", [])
    if isinstance(sources, list):
        sources.append(source)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Import Chengyang's Vietnamese Quốc Ngữ <-> Chữ Nôm workbook. "
            "Blank mappings are intentionally ignored and can be filled later."
        )
    )
    parser.add_argument("xlsx", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/nom_dictionary.json"),
    )
    args = parser.parse_args()

    if not args.xlsx.exists():
        raise SystemExit(f"Excel file not found: {args.xlsx}")

    workbook = load_workbook(args.xlsx, read_only=True, data_only=True)

    for required_sheet in (VOCAB_SHEET, CHAR_SHEET):
        if required_sheet not in workbook.sheetnames:
            raise SystemExit(
                f'Required sheet "{required_sheet}" not found. '
                f"Available: {workbook.sheetnames}"
            )

    raw_entries: list[dict[str, object]] = []
    source_order = 0

    # 1) 词汇:
    # Import only rows where BOTH 喃字 and 国语字 exist.
    vocab = workbook[VOCAB_SHEET]
    columns = header_map(vocab)

    nom_i = require_column(columns, "喃字", VOCAB_SHEET)
    latin_i = require_column(columns, "国语字", VOCAB_SHEET)

    chinese_i = columns.get("中文")
    pos_i = columns.get("词性")
    nom_example_i = columns.get("喃字例句")
    latin_example_i = columns.get("国语字例句")
    chinese_example_i = columns.get("中文例句")

    skipped_blank_vocab = 0

    for row_number, row in enumerate(
        vocab.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        nom_cell = optional_value(row, nom_i)
        latin = optional_value(row, latin_i)

        if not nom_cell or not latin:
            skipped_blank_vocab += 1
            continue

        source_meta: dict[str, object] = {
            "sheet": VOCAB_SHEET,
            "row": row_number,
        }

        chinese = optional_value(row, chinese_i)
        pos = optional_value(row, pos_i)
        nom_example = optional_value(row, nom_example_i)
        latin_example = optional_value(row, latin_example_i)
        chinese_example = optional_value(row, chinese_example_i)

        if chinese:
            source_meta["chinese"] = chinese
        if pos:
            source_meta["partOfSpeech"] = pos
        if nom_example:
            source_meta["nomExample"] = nom_example
        if latin_example:
            source_meta["latinExample"] = latin_example
        if chinese_example:
            source_meta["chineseExample"] = chinese_example

        for nom in split_nom_variants(nom_cell):
            source_order += 1
            raw_entries.append(
                {
                    "latin": latin,
                    "nom": nom,
                    "priority": 0,
                    "metadata": {
                        "sourceOrder": source_order,
                        "sources": [dict(source_meta)],
                    },
                }
            )

    # 2) 喃字:
    # Import character mappings. Multiple Latin readings separated by commas
    # or semicolons become separate candidates.
    chars = workbook[CHAR_SHEET]
    columns = header_map(chars)

    nom_i = require_column(columns, "喃字", CHAR_SHEET)
    latin_i = require_column(columns, "国语字", CHAR_SHEET)
    unicode_i = columns.get("Unicode")

    skipped_blank_char = 0

    for row_number, row in enumerate(
        chars.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        nom = optional_value(row, nom_i)
        latin_cell = optional_value(row, latin_i)

        if not nom or not latin_cell:
            skipped_blank_char += 1
            continue

        unicode_value = optional_value(row, unicode_i)

        for latin in split_latin_variants(latin_cell):
            source_order += 1
            source_meta: dict[str, object] = {
                "sheet": CHAR_SHEET,
                "row": row_number,
            }
            if unicode_value:
                source_meta["unicode"] = unicode_value
            if latin_cell != latin:
                source_meta["rawLatinCell"] = latin_cell

            raw_entries.append(
                {
                    "latin": latin,
                    "nom": nom,
                    "priority": 0,
                    "metadata": {
                        "sourceOrder": source_order,
                        "sources": [source_meta],
                    },
                }
            )

    # Merge only exact duplicate Latin+Nôm pairs.
    # Different Nôm choices for the same Latin text remain separate candidates.
    merged: dict[tuple[str, str], dict[str, object]] = {}

    for entry in raw_entries:
        latin = clean_text(entry["latin"])
        nom = clean_text(entry["nom"])
        key = (normalize_latin(latin), nom)

        existing = merged.get(key)
        if existing is None:
            merged[key] = entry
            continue

        existing_metadata = existing.get("metadata")
        new_metadata = entry.get("metadata")

        if isinstance(existing_metadata, dict) and isinstance(new_metadata, dict):
            new_sources = new_metadata.get("sources")
            if isinstance(new_sources, list):
                for source in new_sources:
                    if isinstance(source, dict):
                        append_source(existing_metadata, source)

    entries = list(merged.values())

    payload = {
        "format": "clyven-nom-dictionary",
        "version": 1,
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFile": args.xlsx.name,
        "sourceSheets": [VOCAB_SHEET, CHAR_SHEET],
        "rules": {
            "ignoreBlankMappings": True,
            "splitLatinSeparators": [",", "，", ";", "；"],
            "splitNomAlternativeStyles": ["A/B", "A（B）"],
        },
        "stats": {
            "rawEntries": len(raw_entries),
            "uniqueEntries": len(entries),
            "duplicatesMerged": len(raw_entries) - len(entries),
            "blankVocabularyRowsIgnored": skipped_blank_vocab,
            "blankCharacterRowsIgnored": skipped_blank_char,
        },
        "entries": entries,
    }

    output = args.output
    if not output.is_absolute():
        output = Path(__file__).resolve().parent.parent / output

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    print()
    print("Clyven Nom dictionary imported.")
    print(f"Source: {args.xlsx}")
    print(f"Vocabulary blank rows ignored: {skipped_blank_vocab}")
    print(f"Character blank rows ignored : {skipped_blank_char}")
    print(f"Raw mapping entries          : {len(raw_entries)}")
    print(f"Unique mapping entries       : {len(entries)}")
    print(f"Duplicates merged            : {len(raw_entries) - len(entries)}")
    print(f"Output                       : {output}")
    print()
    print("Blank mappings were NOT treated as errors.")
    print("Fill them in the Excel later and simply re-import.")


if __name__ == "__main__":
    main()
